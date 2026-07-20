import argparse
import re
import sys
from pathlib import Path
from zipfile import ZipFile

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.infrastructure.docx.validator import validate_docx  # noqa: E402


def _all_text(path: Path) -> str:
    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [
        paragraph.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    ]
    with ZipFile(path) as archive:
        header_footer_xml = "\n".join(
            archive.read(name).decode("utf-8", errors="ignore")
            for name in archive.namelist()
            if name.startswith(("word/header", "word/footer"))
        )
    return "\n".join([*paragraphs, *cells, header_footer_xml])


def main() -> None:
    parser = argparse.ArgumentParser(description="验证生成的周报 DOCX")
    parser.add_argument("path", type=Path)
    parser.add_argument("--expect", default="")
    arguments = parser.parse_args()
    validation = validate_docx(arguments.path, expected_sections=set())
    if not validation.valid:
        raise SystemExit(f"DOCX_INVALID errors={','.join(validation.errors)}")
    text = _all_text(arguments.path)
    if arguments.expect and arguments.expect not in text:
        raise SystemExit(f"DOCX_CONTENT_MISSING text={arguments.expect}")
    unresolved = len(re.findall(r"\{\{\s*[^{}]+?\s*\}\}", text))
    if unresolved:
        raise SystemExit(f"DOCX_UNRESOLVED count={unresolved}")
    print("DOCX_OK sections=all unresolved_placeholders=0")


if __name__ == "__main__":
    main()
